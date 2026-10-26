from __future__ import annotations

import csv
from io import BytesIO, StringIO
from contextlib import nullcontext
import sys
import types
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from src.ui.detailed_risk_export import (
    CSV_COLUMNS,
    component_rows,
    detailed_risk_csv_bytes,
    detailed_risk_excel_bytes,
    report_exposure_summary,
)


PARTS = [
    {
        "id": 927,
        "user_id": "private-user-id",
        "analysis_id": "private-analysis-id",
        "workspace_id": "private-workspace-id",
        "mpn": "LM358N",
        "description": "<p>Dual operational amplifier</p>",
        "manufacturer": "onsemi",
        "quantity": 6,
        "lifecycle_status": "Obsolete",
        "stock_available": 0,
        "stock_data_sources": [],
        "stock_data_complete": False,
        "supplier_count": 2,
        "supplier_coverage": "DigiKey and Mouser checked; count evidence incomplete",
        "primary_supplier": "Mouser",
        "risk_score": 100,
        "risk_level": "High",
        "risk_reasons": ["Saved lifecycle status is obsolete", "No stock available"],
        "lifecycle_evidence": [{"source": "Mouser", "status": "Obsolete", "retrieved_at": "2026-10-11"}],
        "lifecycle_conflict": False,
        "alternate_search_status": "Search incomplete: DigiKey unavailable",
        "unit_price": 0,
    },
    {
        "id": 928,
        "user_id": "private-user-id",
        "analysis_id": "private-analysis-id",
        "workspace_id": "private-workspace-id",
        "mpn": "C0603C104K5RACTU",
        "manufacturer": "KEMET",
        "quantity": 10,
        "lifecycle_status": "Active",
        "stock_available": 0,
        "stock_data_sources": [{"supplier": "DigiKey"}],
        "stock_data_complete": True,
        "lead_time_weeks": 8,
        "lead_time_source": "Mouser",
        "supplier_count": 3,
        "risk_score": 20,
        "risk_level": "Low",
        "risk_reasons": "=HYPERLINK(\"https://example.test\",\"Open\")",
        "lifecycle_evidence": [],
        "lifecycle_conflict": False,
        "alternate_search_status": "Search completed",
        "unit_price": 0.12,
    },
]


def _check_excel_export_has_summary_details_and_high_risk_views_without_private_ids():
    data = detailed_risk_excel_bytes(
        PARTS,
        bom_name="Flight Controller Rev B",
        source_filename="flight-controller.csv",
        analyzed_on="2026-10-11T12:33:31.741757+00:00",
    )
    workbook = load_workbook(BytesIO(data), data_only=False)

    assert workbook.sheetnames == ["Summary", "Component Details", "High Risk Parts", "Alternative Candidates"]
    summary = workbook["Summary"]
    assert summary["A1"].value == "Cadivor BOM Risk Report"
    assert summary["A4"].value == "Project / BOM"
    assert summary["B4"].value == "Flight Controller Rev B"
    assert summary["B5"].value == "flight-controller.csv"
    metrics = {row[0]: row[1] for row in summary.iter_rows(min_row=10, max_row=50, values_only=True)}
    assert metrics["Total components"] == 2
    assert metrics["Parts with matched supplier evidence"] == 2
    assert metrics["Risk assessed"] == 2
    assert metrics["Risk coverage"] == "100%"
    assert "Risk ratings are recorded for all 2 parts" in summary["A7"].value
    assert metrics["High risk"] == 1
    assert metrics["Low risk"] == 1
    assert metrics["Stock not reported"] == 1
    assert metrics["Stock coverage incomplete"] == 1
    assert metrics["Stock coverage unavailable"] == 0

    details = workbook["Component Details"]
    headers = [cell.value for cell in details[1]]
    assert "Evidence source" in headers
    assert "Lifecycle evidence" in headers
    assert "Supplier coverage" in headers
    assert "Supplier product URL" in headers
    assert "Datasheet URL" in headers
    assert "Candidate MPNs" in headers
    assert not {"id", "user_id", "analysis_id", "workspace_id"}.intersection(headers)
    first = dict(zip(headers, [cell.value for cell in details[2]]))
    second = dict(zip(headers, [cell.value for cell in details[3]]))
    assert first["Stock"] == "Not reported"
    assert first["Description"] == "Dual operational amplifier"
    assert first["Evidence source"] == "—"
    assert first["Recorded supplier"] == "Mouser"
    assert first["Stock coverage"] == "Incomplete"
    assert "No stock available" not in first["Why"]
    assert "source-backed stock evidence was not saved" in first["Why"]
    assert "Mouser: Obsolete" in first["Lifecycle evidence"]
    assert second["Stock"] == 0
    assert second["Evidence source"] == "DigiKey"
    why_cell = details.cell(3, headers.index("Why") + 1)
    assert why_cell.data_type == "s"
    assert why_cell.value.startswith("'=HYPERLINK")
    assert workbook["High Risk Parts"].max_row == 2

    exported_values = " ".join(str(cell.value) for sheet in workbook.worksheets for row in sheet for cell in row)
    for private_value in ("private-user-id", "private-analysis-id", "private-workspace-id"):
        assert private_value not in exported_values


def _check_csv_export_matches_stress_free_audit_columns_and_marks_unknown_stock():
    payload = detailed_risk_csv_bytes(PARTS)
    assert payload.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(StringIO(payload.decode("utf-8-sig"))))
    assert tuple(rows[0].keys()) == CSV_COLUMNS
    assert rows[0]["MPN"] == "LM358N"
    assert rows[0]["Stock"] == "Not reported"
    assert rows[0]["Evidence source"] == "—"
    assert rows[1]["Stock"] == "0"
    assert rows[1]["Evidence source"] == "DigiKey"
    assert rows[1]["Why"].startswith("'=")


class DetailedRiskExportTests(unittest.TestCase):
    def test_excel_export_has_customer_facing_summary_and_details(self):
        _check_excel_export_has_summary_details_and_high_risk_views_without_private_ids()

    def test_csv_export_matches_stress_free_audit_columns(self):
        _check_csv_export_matches_stress_free_audit_columns_and_marks_unknown_stock()

    def test_unassessed_parts_are_explained_in_summary_and_empty_high_risk_sheet(self):
        workbook = load_workbook(
            BytesIO(
                detailed_risk_excel_bytes(
                    [
                        {
                            "mpn": "CON-USB",
                            "risk_level": "Needs verification",
                            "risk_score": 0,
                            "supplier_coverage": "No configured supplier returned a matching part record.",
                        }
                    ],
                    bom_name="Connector test",
                )
            ),
            data_only=True,
        )

        summary = workbook["Summary"]
        metrics = {
            row[0]: row[1]
            for row in summary.iter_rows(min_row=10, max_row=50, values_only=True)
            if row[0]
        }
        self.assertIn("0 of 1 parts", summary["A7"].value)
        self.assertIn("not evidence of low risk", summary["A7"].value)
        self.assertIn("These counts can refer to the same part", summary["A7"].value)
        self.assertNotIn("1 were not risk-rated and 1", summary["A7"].value)
        self.assertEqual(metrics["Parts with matched supplier evidence"], 0)
        self.assertEqual(metrics["Risk assessed"], 0)
        self.assertEqual(metrics["Risk coverage"], "0%")
        detail = workbook["Component Details"]
        detail_headers = {cell.value: cell.column_letter for cell in detail[1]}
        self.assertEqual(detail[f'{detail_headers["Risk"]}2'].value, "Not assessed")
        self.assertEqual(detail[f'{detail_headers["Risk score"]}2'].value, "Not assessed")
        self.assertIn(
            "Verify the full manufacturer part number and rerun",
            detail[f'{detail_headers["Why"]}2'].value,
        )
        self.assertEqual(detail[f'{detail_headers["Risk"]}2'].fill.fgColor.rgb[-6:], "E2E8F0")
        self.assertIn("1 of 1 part(s) were not assessed", workbook["High Risk Parts"]["A2"].value)
        self.assertIn("not proof of low risk", workbook["High Risk Parts"]["A2"].value)

    def test_csv_risk_matches_the_saved_report_priority_assessment(self):
        from src.ui.ei_bom_report import _part_rows

        part = {
            "mpn": "TEST-ONLY-LIFECYCLE-PART",
            "risk_level": "Medium",
            "risk_score": 55,
            "lifecycle_status": "Obsolete",
            "supplier_count": 2,
            "stock_data_sources": ["Mouser"],
        }

        report_row = _part_rows([part])[0]
        exported_row = component_rows([part])[0]
        self.assertEqual(report_row["level"], "High")
        self.assertEqual(report_row["score"], 85)
        self.assertEqual(exported_row["Risk"], report_row["level"])
        self.assertEqual(exported_row["Risk score"], report_row["score"])

        csv_row = next(csv.DictReader(StringIO(
            detailed_risk_csv_bytes([part]).decode("utf-8-sig")
        )))
        self.assertEqual(csv_row["Risk"], report_row["level"])
        self.assertEqual(csv_row["Risk score"], str(report_row["score"]))

    def test_zero_score_without_supplier_match_is_not_a_low_risk_rating(self):
        row = component_rows(
            [{
                "mpn": "CON-USB",
                "risk_level": "Needs verification",
                "risk_score": 0,
                "supplier_count": 0,
                "supplier_coverage": "No configured supplier returned a matching part record.",
            }]
        )[0]
        self.assertEqual(row["Risk"], "Not assessed")
        self.assertEqual(row["Risk score"], "Not assessed")
        self.assertIn("stock, lifecycle, and lead-time risk remain unknown", row["Why"])

    def test_related_legacy_catalog_lead_is_explained_without_rating_original(self):
        row = component_rows([{
            "mpn": "PCF8574P",
            "risk_level": "Needs verification",
            "risk_score": 0,
            "supplier_count": 0,
            "alternate_search_status": (
                "No exact supplier match. Legacy route Philips Semiconductors → NXP "
                "found related catalog lead(s) for review only: PCF8574P,112 (DigiKey). "
                "orderable package variant; verify the package and ordering suffix; "
                "no candidate stock, price, or lifecycle is attributed to PCF8574P.; "
                "Candidates returned"
            ),
        }])[0]

        self.assertEqual(row["Risk"], "Not assessed")
        self.assertIn("PCF8574P,112 (DigiKey)", row["Why"])
        self.assertIn("this BOM part remains unassessed", row["Why"])

    def test_cost_exposure_and_longest_lead_time_use_recorded_evidence_only(self):
        parts = [
            {
                "mpn": "LM358N", "description": "Recommended Alt LM358DR2G",
                "quantity": 10, "unit_price": 2.0, "primary_supplier": "Mouser",
                "supplier_offers": [{
                    "mpn": "LM358N", "distributor": "Mouser", "unit_price": 2.0,
                    "currency": "USD",
                }],
                "risk_level": "High", "lifecycle_status": "Obsolete",
                "lead_time_weeks": 3, "lead_time_source": "Mouser",
            },
            {
                "mpn": "1N4148", "quantity": 50, "unit_price": 0.5,
                "primary_supplier": "DigiKey", "supplier_offers": [{
                    "mpn": "1N4148", "distributor": "DigiKey", "unit_price": 0.5,
                    "currency": "EUR",
                }],
                "risk_level": "Low", "lifecycle_status": "Active",
                "lead_time_weeks": 62, "lead_time_source": "DigiKey",
            },
            {
                "mpn": "UNPRICED-CURRENCY", "quantity": 4, "unit_price": 1.25,
                "primary_supplier": "Newark", "risk_level": "High",
                "lifecycle_status": "Active",
            },
        ]
        rows = component_rows(parts)

        self.assertEqual(rows[0]["Extended cost"], 20)
        self.assertEqual(rows[0]["Unit price currency"], "USD")
        self.assertEqual(rows[0]["Description-mentioned alternatives"], "LM358DR2G")
        self.assertEqual(rows[2]["Extended cost"], 5)
        self.assertEqual(rows[2]["Unit price currency"], "Not reported")

        workbook = load_workbook(
            BytesIO(detailed_risk_excel_bytes(parts, bom_name="Exposure test")),
            data_only=True,
        )
        metrics = {
            row[0]: row[1]
            for row in workbook["Summary"].iter_rows(min_row=10, max_row=50, values_only=True)
            if row[0]
        }
        self.assertEqual(metrics["Extended BOM amount by recorded currency"], "EUR 25.00; USD 20.00")
        self.assertEqual(
            metrics["Priced high-risk / lifecycle extended cost by currency"],
            "USD 20.00",
        )
        self.assertEqual(metrics["High-risk / lifecycle rows without extended cost"], 0)
        self.assertEqual(metrics["Priced high-risk rows missing currency"], 1)
        self.assertEqual(metrics["Longest reported part lead time"], "62 weeks — 1N4148 (DigiKey)")
        self.assertEqual(metrics["Components with lead-time data"], 2)
        summary_text = " ".join(str(cell.value) for row in workbook["Summary"] for cell in row if cell.value)
        self.assertIn("not a forecast of project delay", summary_text)
        self.assertIn("not a confirmed or approved cross", summary_text)
        self.assertIn("not total exposure", summary_text)

    def test_candidate_and_source_evidence_are_saved_in_the_export_without_certifying_fit(self):
        parts = [{
            "mpn": "LM317DCYR",
            "description": "Adjustable voltage regulator",
            "manufacturer": "Texas Instruments",
            "risk_level": "Medium",
            "risk_score": 35,
            "risk_reasons": "Reported manufacturer lead time is 26 weeks or longer",
            "lifecycle_status": "Active",
            "product_url": "https://supplier.example/LM317DCYR",
            "datasheet_url": "https://ti.example/lm317.pdf",
            "sourcing_updated_at": "2026-10-21T12:00:00Z",
            "alternate_search_status": "Candidates returned",
            "alternate_candidates": [{
                "Alternative Part": "LM317BS-TR",
                "Manufacturer": "STMicroelectronics",
                "Supplier": "DigiKey",
                "Classification": "Supplier-listed similar",
                "Evidence Type": "Catalog search",
                "package": "SOT-223",
                "pin_count": 3,
                "Lifecycle": "Active",
                "Product URL": "https://supplier.example/LM317BS-TR",
                "Datasheet URL": "https://st.example/lm317bs.pdf",
                "Retrieved At": "2026-10-21T12:00:00Z",
            }],
        }]

        rows = component_rows(parts)
        self.assertEqual(rows[0]["Candidate MPNs"], "LM317BS-TR")
        self.assertEqual(rows[0]["Supplier product URL"], "https://supplier.example/LM317DCYR")
        self.assertEqual(rows[0]["Datasheet URL"], "https://ti.example/lm317.pdf")

        workbook = load_workbook(BytesIO(detailed_risk_excel_bytes(parts, bom_name="Regulator")))
        self.assertEqual(workbook.sheetnames[-1], "Alternative Candidates")
        sheet = workbook["Alternative Candidates"]
        headers = [cell.value for cell in sheet[4]]
        candidate = dict(zip(headers, [cell.value for cell in sheet[5]]))
        self.assertEqual(candidate["Source MPN"], "LM317DCYR")
        self.assertEqual(candidate["Candidate MPN"], "LM317BS-TR")
        self.assertEqual(candidate["Relationship label"], "Supplier-listed similar")
        self.assertIn("not established", candidate["Approval status"])
        self.assertEqual(sheet.cell(5, headers.index("Supplier product URL") + 1).hyperlink.target,
                         "https://supplier.example/LM317BS-TR")
        self.assertEqual(workbook["Component Details"].cell(2, 8).hyperlink.target,
                         "https://supplier.example/LM317DCYR")
        summary_text = " ".join(
            str(cell.value) for row in workbook["Summary"] for cell in row if cell.value
        )
        self.assertIn("26+ weeks +30", summary_text)
        self.assertIn("60–100 High", summary_text)
        workbook.close()

        csv_rows = list(csv.DictReader(StringIO(
            detailed_risk_csv_bytes(parts).decode("utf-8-sig")
        )))
        self.assertEqual(csv_rows[0]["Candidate MPNs"], "LM317BS-TR")
        self.assertEqual(csv_rows[0]["Risk score"], "35")
        self.assertEqual(csv_rows[0]["Datasheet URL"], "https://ti.example/lm317.pdf")

    def test_saved_unit_price_currency_drives_financial_rollup(self):
        parts = [{
            "mpn": "BC547B",
            "quantity": 150,
            "unit_price": 0.13,
            "unit_price_currency": "USD",
            "primary_supplier": "DigiKey",
            "risk_level": "Low",
            "lifecycle_status": "Active",
        }]
        rows = component_rows(parts)
        self.assertEqual(rows[0]["Unit price currency"], "USD")
        self.assertEqual(rows[0]["Extended cost"], 19.5)

        workbook = load_workbook(
            BytesIO(detailed_risk_excel_bytes(parts, bom_name="Currency rollup test")),
            data_only=True,
        )
        metrics = {
            row[0]: row[1]
            for row in workbook["Summary"].iter_rows(min_row=10, max_row=50, values_only=True)
            if row[0]
        }
        self.assertEqual(metrics["Extended BOM amount by recorded currency"], "USD 19.50")
        self.assertEqual(metrics["Priced components missing currency"], 0)

    def test_approved_alternatives_are_included_in_both_downloads(self):
        approval = {
            "original_part": "LM358N",
            "alternative_part": "LM358DR2G",
            "decision": "Approved",
            "project_name": "Flight Controller Rev B",
            "engineer_name": "Engineer",
            "engineering_note": "Reviewed against the datasheet",
        }
        excel = load_workbook(
            BytesIO(
                detailed_risk_excel_bytes(
                    PARTS,
                    bom_name="Flight Controller Rev B",
                    approved_alternatives=[approval],
                )
            ),
            data_only=True,
        )
        details = excel["Component Details"]
        headers = [cell.value for cell in details[1]]
        first = dict(zip(headers, [cell.value for cell in details[2]]))
        self.assertEqual(first["Approved alternatives"], "LM358DR2G")

        rows = list(
            csv.DictReader(
                StringIO(
                    detailed_risk_csv_bytes(PARTS, [approval]).decode("utf-8-sig")
                )
            )
        )
        self.assertEqual(rows[0]["Approved alternatives"], "LM358DR2G")

    def test_approved_alternative_updates_mitigation_and_open_cost_without_erasing_source_risk(self):
        parts = [{
            "mpn": "MC68HC11E1CFNE2",
            "quantity": 4,
            "unit_price": 3.5,
            "unit_price_currency": "USD",
            "risk_level": "High",
            "risk_score": 88,
            "lifecycle_status": "Obsolete",
        }]
        approval = {
            "original_part": "MC68HC11E1CFNE2",
            "alternative_part": "MC68HC11E1CFNE3",
            "decision": "Approved",
            "engineer_name": "Engineer",
            "project_name": "Controller",
        }

        row = component_rows(parts, [approval])[0]
        self.assertEqual(row["Risk"], "High")
        self.assertEqual(row["Risk score"], 88)
        self.assertEqual(row["Mitigation status"], "Approved alternative (BOM update pending)")
        self.assertEqual(row["Approved alternatives"], "MC68HC11E1CFNE3")

        exposure = report_exposure_summary([row])
        self.assertEqual(exposure["gross_high_risk_amount_by_currency"], "USD 14.00")
        self.assertEqual(exposure["approved_high_risk_amount_by_currency"], "USD 14.00")
        self.assertEqual(exposure["high_risk_amount_by_currency"], "Not available from recorded currency data")

        exported = list(csv.DictReader(StringIO(
            detailed_risk_csv_bytes(parts, [approval]).decode("utf-8-sig")
        )))[0]
        self.assertEqual(exported["Mitigation status"], "Approved alternative (BOM update pending)")
        self.assertEqual(exported["Approved alternatives"], "MC68HC11E1CFNE3")

        workbook = load_workbook(BytesIO(detailed_risk_excel_bytes(
            parts,
            bom_name="Controller",
            approved_alternatives=[approval],
        )), data_only=True)
        summary = {
            values[0]: values[1]
            for values in workbook["Summary"].iter_rows(min_row=10, max_row=60, values_only=True)
            if values[0]
        }
        details = workbook["Component Details"]
        headers = [cell.value for cell in details[1]]
        saved = dict(zip(headers, [cell.value for cell in details[2]]))
        self.assertEqual(summary["Priced high-risk / lifecycle extended cost by currency"], "Not available from recorded currency data")
        self.assertEqual(summary["Gross priced high-risk / lifecycle extended cost by currency"], "USD 14.00")
        self.assertEqual(saved["Risk"], "High")
        self.assertEqual(saved["Mitigation status"], "Approved alternative (BOM update pending)")
        self.assertEqual(saved["Approved alternatives"], "MC68HC11E1CFNE3")

    def test_detailed_risk_page_delivers_the_customer_exports(self):
        from src.ui import ei_bom_report

        downloads = []
        fake_streamlit = types.ModuleType("streamlit")
        fake_streamlit.session_state = {"cadivor_show_detailed_risk": True}
        fake_streamlit.markdown = lambda *_args, **_kwargs: None
        fake_streamlit.html = lambda *_args, **_kwargs: None
        fake_streamlit.container = lambda **_kwargs: nullcontext()
        fake_streamlit.button = lambda *_args, **_kwargs: False
        fake_streamlit.columns = lambda count: [nullcontext() for _ in range(count)]
        fake_streamlit.download_button = lambda label, **kwargs: downloads.append((label, kwargs))

        with patch.dict(sys.modules, {"streamlit": fake_streamlit}), patch.object(
            ei_bom_report, "render_detailed_risk_rows"
        ):
            ei_bom_report.render_engineering_intelligence_report(
                analysis={
                    "id": "private-analysis-id",
                    "project_name": "Flight Controller Rev B",
                    "filename": "flight-controller.csv",
                    "created_at": "2026-10-11T12:33:00+00:00",
                },
                parts=PARTS,
                alternatives=[],
                health_score=72,
            )

        self.assertEqual([item[0] for item in downloads], ["Download Excel Report", "Download CSV Audit"])
        self.assertEqual(downloads[0][1]["file_name"], "cadivor-detailed-risk-report.xlsx")
        workbook = load_workbook(BytesIO(downloads[0][1]["data"]), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Summary", "Component Details", "High Risk Parts", "Alternative Candidates"])
        self.assertEqual(downloads[1][1]["file_name"], "cadivor-detailed-risk-report.csv")
        self.assertTrue(downloads[1][1]["data"].startswith(b"\xef\xbb\xbf"))

