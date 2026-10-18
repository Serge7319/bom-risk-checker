"""Security and response-boundary checks for the anonymous audit."""

from __future__ import annotations

import hashlib
import hmac
import importlib.util
import sys
import types
import unittest
from contextlib import nullcontext
from pathlib import Path
from unittest.mock import patch

from src import public_bom_stress_test as funnel


class PublicBomStressTest(unittest.TestCase):
    def test_free_audit_opens_in_branded_page_below_banner(self):
        root = Path(__file__).resolve().parents[1]
        homepage = (root / "marketing-web/index.html").read_text()
        script = (root / "marketing-web/app.js").read_text()
        entrypoint = (root / "streamlit_app.py").read_text()
        self.assertIn('id="bomStressTestFrame"', homepage)
        self.assertIn('id="bomStressTestSection"', homepage)
        self.assertNotIn('id="bomStressTestFullPage"', homepage)
        self.assertIn('aria-labelledby="stressFunnelTitle" hidden', homepage)
        self.assertIn("CADIVOR_PUBLIC_BOM_STRESS_TEST_ENABLED === true", script)
        self.assertIn("window.CADIVOR_PUBLIC_BOM_STRESS_TEST_ENABLED = true;", homepage)
        self.assertEqual(homepage.count('data-stress-cta="true"'), 3)
        self.assertLess(homepage.index('class="hero experience-scene"'), homepage.index('data-page="analyze"'))
        self.assertLess(homepage.index('class="page-hero stress-page-hero"'), homepage.index('class="stress-funnel"'))
        self.assertIn("a.href = '#/analyze'", script)
        self.assertNotIn("stressFullPage", script)
        self.assertIn("if (page === 'analyze')", script)
        self.assertIn("new URLSearchParams({ public: 'stress', embed: 'true' })", script)
        self.assertIn("params.set(key, value.slice(0, 80))", script)
        self.assertLess(entrypoint.index('st.query_params.get("public"'),
                        entrypoint.index("ensure_authenticated_or_stop()"))

    def test_direct_audit_keeps_brand_without_distraction_links_or_embedded_header(self):
        rendered = []
        fake_st = types.ModuleType("streamlit")
        fake_st.query_params = {}
        fake_st.markdown = lambda html, **kwargs: rendered.append(html)
        fake_st.container = lambda **kwargs: nullcontext()
        fake_st.info = lambda *_: None
        filename = Path(__file__).resolve().parents[1] / "src/public_bom_stress_ui.py"
        spec = importlib.util.spec_from_file_location("stress_public_header_test", filename)
        module = importlib.util.module_from_spec(spec)
        with patch.dict(sys.modules, {"streamlit": fake_st}):
            spec.loader.exec_module(module)
            with patch.object(module, "enabled", return_value=False):
                module.render_public_bom_stress_test()
                self.assertIn('class="cv-stress-site-header"', "".join(rendered))
                self.assertIn('href="https://www.cadivor.com/#/home"', "".join(rendered))
                self.assertIn('class="cv-stress-header-note">Free BOM audit', "".join(rendered))
                self.assertNotIn('href="https://www.cadivor.com/#/product"', "".join(rendered))
                self.assertNotIn('href="https://www.cadivor.com/#/pricing"', "".join(rendered))
                self.assertNotIn('href="https://app.cadivor.com/?auth=login"', "".join(rendered))
                rendered.clear()
                fake_st.query_params = {"embed": "true"}
                module.render_public_bom_stress_test()
                self.assertNotIn('class="cv-stress-site-header"', "".join(rendered))

    def test_daily_limit_does_not_offer_a_retry_that_cannot_succeed(self):
        captions, buttons, errors = [], [], []
        fake_st = types.ModuleType("streamlit")
        fake_st.query_params = {"embed": "true"}
        fake_st.session_state = {"cv_stress_error": funnel.PUBLIC_BOM_LIMIT_MESSAGE}
        fake_st.markdown = lambda *_args, **_kwargs: None
        fake_st.container = lambda **_kwargs: nullcontext()
        fake_st.write = lambda *_args: None
        fake_st.caption = lambda message: captions.append(message)
        fake_st.error = lambda message: errors.append(message)
        fake_st.button = lambda label, **_kwargs: buttons.append(label) or False
        fake_st.file_uploader = lambda *_args, **_kwargs: None
        filename = Path(__file__).resolve().parents[1] / "src/public_bom_stress_ui.py"
        spec = importlib.util.spec_from_file_location("stress_rate_limit_ui_test", filename)
        module = importlib.util.module_from_spec(spec)
        with patch.dict(sys.modules, {"streamlit": fake_st}):
            spec.loader.exec_module(module)
            with patch.object(module, "enabled", return_value=True):
                module.render_public_bom_stress_test()
        self.assertEqual(errors, [funnel.PUBLIC_BOM_LIMIT_MESSAGE])
        self.assertTrue(any("rolling 24-hour window" in caption for caption in captions))
        self.assertNotIn("Retry audit", buttons)

    def test_visitor_must_be_signed_by_ingress(self):
        secret = "a" * 48
        stamp = 1_000_000
        ip = "2001:db8::1"
        digest = hmac.new(secret.encode(), f"v1|{stamp}|{ip}".encode(), hashlib.sha256).hexdigest()
        header = f"v1;{stamp};{ip};{digest}"
        self.assertEqual(len(funnel.visitor_hash(header, now=stamp, secret=secret)), 64)
        with self.assertRaises(funnel.StressTestError):
            funnel.visitor_hash(header.replace(ip, "2001:db8::2"), now=stamp, secret=secret)
        with self.assertRaises(funnel.StressTestError):
            funnel.visitor_hash(header, now=stamp + 7201, secret=secret)
        with self.assertRaises(funnel.StressTestError):
            funnel.visitor_hash("1.2.3.4", now=stamp, secret=secret)

    def test_email_validation_accepts_consumer_aliases_and_bounded_parser(self):
        for address in ("person@gmail.com", "person+cadivor-test@gmail.com", "person@outlook.com"):
            self.assertEqual(funnel.normalize_email_address(address), address)
        for address in ("bad-address", "person@localhost"):
            with self.subTest(address=address), self.assertRaises(funnel.StressTestError):
                funnel.normalize_email_address(address)
        self.assertEqual(
            funnel.normalize_email_address(" Director@Acme-Systems.com "),
            "director@acme-systems.com",
        )
        bom = b"Manufacturer Part Number,Qty\nMCP2551-I/SN,2\nTPS5430DDAR,5\n"
        self.assertEqual(funnel.parse_bom("board.csv", bom)[0], {"mpn": "MCP2551-I/SN", "quantity": 2})
        with self.assertRaises(funnel.StressTestError):
            funnel.parse_bom("board.csv", b"MPN,Qty\n=HYPERLINK(\"bad\"),1\n")
        with self.assertRaises(funnel.StressTestError):
            funnel.parse_bom("board.csv", b"MPN,Qty\n" + b"PART,1\n" * 31)
        with self.assertRaises(funnel.StressTestError):
            funnel.parse_bom("board.csv", b"x" * (funnel.MAX_BYTES + 1))
        self.assertEqual(funnel._safe_export_cell(" =WEBSERVICE(\"https://example.com\")"),
                         "' =WEBSERVICE(\"https://example.com\")")

    def test_release_readiness_campaign_attribution_is_bounded_and_saved_on_lead(self):
        params = {
            "utm_source": "release-readiness-guide",
            "utm_medium": "owned content",
            "utm_campaign": "release-readiness",
        }
        source = funnel.lead_source_from_query(params)
        self.assertEqual(source, "utm:release-readiness-guide:owned-content:release-readiness")
        self.assertEqual(funnel.lead_source_from_query({}), "homepage_stress_test")
        self.assertNotIn("@", funnel.lead_source_from_query({"utm_source": "hello@company.com"}))

        class FakeSupabase:
            def __init__(self):
                self.current_table = ""
                self.lead = None
                self.is_update = False

            def table(self, name):
                self.current_table = name
                self.is_update = False
                return self

            def select(self, *_):
                return self

            def eq(self, *_):
                return self

            def single(self):
                return self

            def update(self, *_):
                self.is_update = True
                return self

            def upsert(self, payload, **_):
                self.lead = payload
                return self

            def execute(self):
                if self.current_table == "cadivor_public_bom_stress_tests":
                    data = [{"id": "report-1"}] if self.is_update else {
                        "id": "report-1", "row_count": 10,
                        "status": "ready", "work_email": None,
                        "results": [],
                    }
                    return types.SimpleNamespace(data=data)
                return types.SimpleNamespace(data=None)

        client = FakeSupabase()
        with patch.object(funnel, "_service_client", return_value=client):
            funnel.capture_lead("report-1", "engineer@company.com", source=source)
        self.assertEqual(client.lead["source"], source)

    def test_release_readiness_guide_is_linked_from_marketing_entry_points(self):
        root = Path(__file__).resolve().parents[1]
        homepage = (root / "marketing-web/index.html").read_text()
        guide = (root / "marketing-web/release-readiness.html").read_text()
        script = (root / "marketing-web/app.js").read_text()
        self.assertIn("Know when your hardware is", homepage)
        self.assertIn("Cadivor — Hardware Release Readiness", script)
        self.assertIn('href="/release-readiness.html"', homepage)
        self.assertIn("7 BOM Checks", guide)
        self.assertEqual(guide.count('id="check-'), 7)
        self.assertIn("buildPublicStressUrl()", script)
        self.assertIn("utm_campaign", script)

    def test_rate_slot_is_reserved_before_any_provider_call_and_locked_rows_never_return(self):
        class FakeUpdate:
            def __init__(self):
                self.saved = None

            def table(self, name):
                self.table_name = name
                return self

            def update(self, payload):
                self.saved = payload
                return self

            def eq(self, *_):
                return self

            def execute(self):
                return None

        fake = FakeUpdate()
        calls = []

        def reserve(_):
            calls.append("reserve")
            return "00000000-0000-0000-0000-000000000000"

        def enrich(index, row):
            self.assertEqual(calls[0], "reserve")
            return index, {
                **row, "risk": "High" if index >= 5 else "Low",
                "verified": True, "risk_reasons": [f"private-evidence-{index}"],
                "lifecycle": "Active", "stock": 7, "lead_time_weeks": 1.0,
                "supplier_count": 2, "sources": "DigiKey, Mouser",
            }

        csv = ("MPN,Qty\n" + "".join(f"PART{i},1\n" for i in range(10))).encode()
        with patch.object(funnel, "enabled", return_value=True), \
             patch.object(funnel, "visitor_hash", return_value="f" * 64), \
             patch.object(funnel, "_reserve", side_effect=reserve), \
             patch.object(funnel, "_service_client", return_value=fake), \
             patch.object(funnel, "_enrich_one", side_effect=enrich):
            teaser = funnel.run_audit("prototype.csv", csv, "signed")
        self.assertEqual(len(teaser["preview"]), 5)
        self.assertEqual(teaser["remaining_high_risk"], 5)
        self.assertEqual(len(fake.saved["results"]), 10)
        self.assertNotIn("PART5", str(teaser))
        self.assertNotIn("private-evidence-5", str(teaser))

    def test_full_results_require_a_confirmed_supabase_identity_with_matching_email(self):
        class FakeQuery:
            def __init__(self):
                self.filters = []

            def table(self, name):
                self.name = name
                return self

            def select(self, *_):
                return self

            def eq(self, key, value):
                self.filters.append((key, value))
                return self

            def gte(self, *_):
                return self

            def order(self, *_args, **_kwargs):
                return self

            def limit(self, *_):
                return self

            def execute(self):
                return types.SimpleNamespace(data=[{"results": [{"mpn": "PRIVATE-PART"}]}])

        query = FakeQuery()
        user = types.SimpleNamespace(email="engineer@company.com", email_confirmed_at=None)
        auth = types.SimpleNamespace(auth=types.SimpleNamespace(
            get_user=lambda token: types.SimpleNamespace(user=user)))
        supabase = types.SimpleNamespace(create_client=lambda *_: auth)
        with patch.dict(sys.modules, {"supabase": supabase}), \
             patch.object(funnel, "enabled", return_value=True), \
             patch.object(funnel, "get_secret", return_value="configured"), \
             patch.object(funnel, "_service_client", return_value=query):
            self.assertIsNone(funnel.verified_report(""))
            self.assertIsNone(funnel.verified_report("unverified-jwt"))
            self.assertEqual(query.filters, [])
            user.email_confirmed_at = "2026-10-03T10:00:00Z"
            report = funnel.verified_report("confirmed-jwt")
        self.assertEqual(report["results"][0]["mpn"], "PRIVATE-PART")
        self.assertIn(("work_email", "engineer@company.com"), query.filters)

    def test_supplier_outage_is_unverified_not_a_fabricated_high_risk_dependency(self):
        supplier_module = types.SimpleNamespace(get_best_part_data=lambda _: {
            "supplier_data_verified": False, "stock_total": 0,
            "lifecycle_status": "Unknown", "supplier_count": 0,
        })
        with patch.dict(sys.modules, {"integrations.supplier_aggregator": supplier_module}):
            _, result = funnel._enrich_one(0, {"mpn": "PART1", "quantity": 2})
        self.assertFalse(result["verified"])
        self.assertEqual(result["risk"], "Needs verification")
        self.assertIsNone(result["stock"])

    def test_email_link_uses_existing_verified_callback_without_password(self):
        sent = []
        auth = types.SimpleNamespace(auth=types.SimpleNamespace(
            sign_in_with_otp=lambda payload: sent.append(payload)))
        supabase = types.SimpleNamespace(create_client=lambda *_: auth)
        callback = types.SimpleNamespace(
            signup_confirmation_redirect_url=lambda: "https://app.cadivor.com/?cadivor_signup_confirm=1")
        with patch.dict(sys.modules, {
            "supabase": supabase, "src.auth_signup_confirmation": callback,
        }), patch.object(funnel, "get_secret", return_value="configured"):
            funnel.send_report_verification("Director@Company.com")
        self.assertEqual(sent[0]["email"], "director@company.com")
        self.assertTrue(sent[0]["options"]["should_create_user"])
        self.assertIn("cadivor_signup_confirm=1", sent[0]["options"]["email_redirect_to"])

    def test_verified_email_callback_hands_off_to_reports_only_after_success(self):
        fake_st = types.ModuleType("streamlit")
        fake_st.session_state = {}
        fake_st.query_params = {
            "cadivor_signup_confirm": "1", "cadivor_stress_report": "1",
            "token_hash": "one-time-hash", "type": "email",
        }
        auth_state = types.ModuleType("src.auth_state")
        for name in ("APP_LOGIN", "APP_SIGNUP", "APP_SIGNUP_CONFIRMATION_LINK_READY",
                     "APP_SIGNUP_CONFIRMATION_INVALID",
                     "APP_SIGNUP_CONFIRMATION_SUCCESS", "AUTH_SIGNED_OUT",
                     "SIGNUP_PENDING_EMAIL_KEY"):
            setattr(auth_state, name, name)
        auth_state.log_auth_diagnostic = lambda *args, **kwargs: None
        filename = Path(__file__).resolve().parents[1] / "src/auth_signup_confirmation.py"
        spec = importlib.util.spec_from_file_location("stress_confirmation_test", filename)
        module = importlib.util.module_from_spec(spec)
        with patch.dict(sys.modules, {"streamlit": fake_st, "src.auth_state": auth_state}):
            spec.loader.exec_module(module)
            with patch.object(module, "_activate_from_token_hash", return_value=module.RESULT_SESSION_READY) as activate:
                module.apply_signup_confirmation_from_query(object())
                activate.assert_not_called()
                self.assertTrue(module.signup_confirmation_verification_pending())
                module.verify_pending_signup_confirmation(object())
                activate.assert_called_once()
            self.assertTrue(fake_st.session_state["cadivor_stress_landing_pending"])
            self.assertNotIn("token_hash", fake_st.query_params)
            self.assertNotIn("cadivor_stress_report", fake_st.query_params)
            fake_st.session_state.clear()
            fake_st.query_params.update({
                "cadivor_signup_confirm": "1", "cadivor_stress_report": "1",
                "token_hash": "bad-hash", "type": "email",
            })
            with patch.object(module, "_activate_from_token_hash", return_value=module.RESULT_INVALID):
                module.apply_signup_confirmation_from_query(object())
            self.assertNotIn("cadivor_stress_landing_pending", fake_st.session_state)

    def test_verified_handoff_keeps_reports_in_url_and_session(self):
        fake_st = types.ModuleType("streamlit")
        fake_st.session_state = {"cadivor_stress_landing_pending": True}
        fake_st.query_params = {}
        components = types.ModuleType("streamlit.components")
        component_v1 = types.ModuleType("streamlit.components.v1")
        components.v1 = component_v1
        fake_st.components = components
        auth_diagnostics = types.ModuleType("src.auth_diagnostics")
        auth_diagnostics.log_auth_correlation = lambda *args, **kwargs: None
        auth_cookies = types.ModuleType("src.auth_cookies")
        auth_cookies.persist_session_auth_cookie = lambda *args, **kwargs: None
        auth_cookies.get_auth_cookie_manager = lambda **kwargs: object()
        filename = Path(__file__).resolve().parents[1] / "src/auth_state.py"
        spec = importlib.util.spec_from_file_location("stress_auth_state_test", filename)
        module = importlib.util.module_from_spec(spec)
        with patch.dict(sys.modules, {
            "streamlit": fake_st, "streamlit.components": components,
            "streamlit.components.v1": component_v1,
            "src.auth_diagnostics": auth_diagnostics, "src.auth_cookies": auth_cookies,
        }):
            spec.loader.exec_module(module)
            session = types.SimpleNamespace(access_token="token", refresh_token="refresh")
            module.mark_authenticated(types.SimpleNamespace(id="user1"), session, object())
        self.assertEqual(fake_st.session_state["cadivor_route"], "Reports")
        self.assertEqual(fake_st.query_params["page"], "Reports")

    def test_audit_downloads_keep_evidence_and_use_readable_numeric_values(self):
        from io import BytesIO, StringIO

        from openpyxl import load_workbook

        row = {
            "mpn": "ATMEGA328P-PU", "matched_mpn": "ATMEGA328P-PU",
            "manufacturer": "Microchip Technology", "quantity": 50,
            "lifecycle": "Conflicting supplier data",
            "lifecycle_evidence_summary": "DigiKey: Active; Mouser: Not Recommended for New Designs",
            "stock": 41842, "stock_evidence": "Reported by DigiKey, Mouser",
            "lead_time_weeks": 27.0, "lead_time_source": "DigiKey",
            "supplier_count": 3, "sources": "DigiKey, Mouser, Newark",
            "supplier_coverage": "3 of 3 configured supplier sources responded successfully.",
            "checked_at": "2026-10-18T23:14:27+00:00", "risk": "High",
            "risk_reasons": ["Supplier lifecycle records conflict"],
        }
        csv_bytes = funnel.build_audit_csv([row])
        self.assertTrue(csv_bytes.startswith(b"\xef\xbb\xbf"))
        self.assertNotIn(b"41842.0", csv_bytes)
        csv_rows = list(__import__("csv").DictReader(StringIO(csv_bytes.decode("utf-8-sig"))))
        self.assertEqual(csv_rows[0]["Reported stock"], "41842")
        self.assertEqual(csv_rows[0]["Matching suppliers"], "3")
        self.assertEqual(csv_rows[0]["Why"], "Supplier lifecycle records conflict")
        self.assertEqual(len(csv_rows[0]), 16)

        preview = funnel.build_audit_preview([row])
        self.assertEqual(list(preview.columns), [
            "BOM MPN", "Manufacturer", "Qty", "Lifecycle", "Reported stock",
            "Lead time (weeks)", "Matching suppliers", "Risk",
        ])
        report = {
            "filename": "enterprise_test_bom_25.csv",
            "created_at": "2026-10-18T23:20:00+00:00",
            "results": [row],
        }
        workbook = load_workbook(BytesIO(funnel.build_audit_xlsx(report)), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Summary", "Component evidence"])
        self.assertEqual(workbook["Summary"]["A1"].value, "Cadivor | BOM Stress Test")
        self.assertEqual(workbook["Summary"]["A6"].value, 1)
        details = workbook["Component evidence"]
        self.assertEqual(details.freeze_panes, "A5")
        self.assertEqual(details.tables["BOMComponentEvidence"].ref, "A4:P5")
        self.assertEqual(details["G5"].value, 41842)
        self.assertEqual(details["K5"].value, 3)
        self.assertEqual(details["P5"].value, "Supplier lifecycle records conflict")

    def test_audit_workbook_escapes_formula_like_supplier_evidence(self):
        from io import BytesIO

        from openpyxl import load_workbook

        row = {
            "mpn": "PART-1", "quantity": 1, "lifecycle": "Active", "risk": "Low",
            "risk_reasons": ['=HYPERLINK("https://example.com")'],
        }
        workbook = load_workbook(BytesIO(funnel.build_audit_xlsx({"results": [row]})))
        cell = workbook["Component evidence"]["P5"]
        self.assertEqual(cell.data_type, "s")
        self.assertTrue(cell.value.startswith("'=HYPERLINK"))

    def test_verified_report_shows_compact_table_and_both_download_formats(self):
        import time

        captured = {"buttons": [], "frames": []}
        fake_st = types.ModuleType("streamlit")
        fake_st.session_state = {
            "user": {"id": "user-1"},
            "cadivor_stress_verified_report_cache": {
                "user_id": "user-1", "at": time.time(),
                "report": {
                    "filename": "sample.csv", "created_at": "2026-10-18T23:20:00+00:00",
                    "results": [{
                        "mpn": "ATMEGA328P-PU", "matched_mpn": "ATMEGA328P-PU",
                        "manufacturer": "Microchip Technology", "quantity": 50,
                        "lifecycle": "Active", "stock": 41842, "stock_evidence": "DigiKey",
                        "lead_time_weeks": 2.0, "lead_time_source": "DigiKey", "supplier_count": 3,
                        "sources": "DigiKey, Mouser, Newark", "supplier_coverage": "3 of 3 sources",
                        "checked_at": "2026-10-18T23:14:27+00:00", "risk": "Low",
                        "risk_reasons": [],
                    }],
                },
            },
        }
        fake_st.markdown = lambda *_args, **_kwargs: None
        fake_st.text = lambda *_args, **_kwargs: None
        fake_st.caption = lambda *_args, **_kwargs: None
        fake_st.info = lambda *_args, **_kwargs: None
        fake_st.dataframe = lambda frame, **_kwargs: captured["frames"].append(frame)
        fake_st.expander = lambda *_args, **_kwargs: nullcontext()
        fake_st.container = lambda *_args, **_kwargs: nullcontext()
        fake_st.selectbox = lambda _label, options, **_kwargs: options[0]
        fake_st.columns = lambda spec, **_kwargs: [nullcontext() for _ in range(spec if isinstance(spec, int) else len(spec))]
        fake_st.write = lambda *_args, **_kwargs: None
        fake_st.download_button = lambda label, **kwargs: captured["buttons"].append((label, kwargs["data"]))
        fake_st.column_config = types.SimpleNamespace(
            TextColumn=lambda *args, **kwargs: (args, kwargs),
            NumberColumn=lambda *args, **kwargs: (args, kwargs),
        )
        with patch.dict(sys.modules, {"streamlit": fake_st}), patch.object(funnel, "enabled", return_value=True):
            funnel.render_verified_report("verified-token")

        self.assertEqual(list(captured["frames"][0].columns), [
            "BOM MPN", "Manufacturer", "Qty", "Lifecycle", "Reported stock",
            "Lead time (weeks)", "Matching suppliers", "Risk",
        ])
        self.assertEqual([label for label, _ in captured["buttons"]], [
            "Download Excel report", "Download CSV data",
        ])
        self.assertTrue(captured["buttons"][0][1].startswith(b"PK"))
        self.assertTrue(captured["buttons"][1][1].startswith(b"\xef\xbb\xbf"))


if __name__ == "__main__":
    unittest.main()
