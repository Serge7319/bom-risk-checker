"""Checkout and credit safety tests; no live Stripe or Supabase credentials."""

from __future__ import annotations

import sys
import types
import unittest
import importlib.util
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from src import one_time_bom
from src.report_generator import save_results_to_excel


USER = "0453a98c-4dd1-4c2d-8435-e4a562de0a50"
ORDER = "0ff5f8f3-8fe9-4d67-aeae-41686fc285f7"
ANALYSIS = "3c3d8352-86c3-49c6-99b5-39c8e949b5a8"


class OneTimeBOMTests(unittest.TestCase):
    def setUp(self):
        self.price_id = "price_test_single_bom"

    def test_feature_stays_off_without_price_and_service_key(self):
        with patch.object(one_time_bom, "get_secret_bool", return_value=True), \
             patch.object(one_time_bom, "get_secret", side_effect=lambda key: {
                 "STRIPE_ONE_TIME_BOM_REPORT_PRICE_ID": "",
                 "SUPABASE_SERVICE_ROLE_KEY": "secret",
             }.get(key)):
            self.assertFalse(one_time_bom.enabled())

    def test_price_is_real_active_nonrecurring_stripe_price(self):
        one_time_bom.price_label.cache_clear()
        fake_stripe = types.SimpleNamespace(Price=types.SimpleNamespace(
            retrieve=Mock(return_value=types.SimpleNamespace(
                active=True, recurring=None, currency="usd", unit_amount=4900,
            )),
        ))
        fake_helper = types.ModuleType("src.stripe_helper")
        fake_helper._ensure_stripe_api_key = Mock()
        with patch.dict(sys.modules, {"stripe": fake_stripe, "src.stripe_helper": fake_helper}):
            self.assertEqual(one_time_bom.price_label(self.price_id), "USD 49.00")
            fake_stripe.Price.retrieve.return_value.recurring = {"interval": "month"}
            with self.assertRaises(one_time_bom.OneTimeBOMError):
                one_time_bom.price_label("price_recurring")
        one_time_bom.price_label.cache_clear()

    def test_checkout_order_is_stored_before_url_is_returned(self):
        client = Mock()
        chain = client.table.return_value
        chain.insert.return_value.execute.return_value.data = [{"id": ORDER}]
        chain.update.return_value.eq.return_value.eq.return_value.eq.return_value.select.return_value.execute.return_value.data = [{"id": ORDER}]
        create = Mock(return_value=types.SimpleNamespace(id="cs_test_one", url="https://checkout.stripe.test/one"))
        fake_helper = types.ModuleType("src.stripe_helper")
        fake_helper.create_one_time_bom_checkout = create
        with patch.dict(sys.modules, {"src.stripe_helper": fake_helper}), \
             patch.object(one_time_bom, "_service_client", return_value=client), \
             patch.object(one_time_bom, "enabled", return_value=True), \
             patch.object(one_time_bom, "pending_checkout", return_value=None), \
             patch.object(one_time_bom, "price_label", return_value="USD 49.00"), \
             patch.object(one_time_bom, "get_secret", return_value=self.price_id):
            url = one_time_bom.begin_checkout(USER, "eng@example.com", "https://app/success", "https://app/cancel")
        self.assertEqual(url, "https://checkout.stripe.test/one")
        create.assert_called_once()
        self.assertEqual(create.call_args.kwargs["order_id"], ORDER)
        chain.update.assert_called_once_with({"stripe_session_id": "cs_test_one"})

    def test_an_open_checkout_cannot_create_a_second_order(self):
        client = Mock()
        with patch.object(one_time_bom, "enabled", return_value=True), \
             patch.object(one_time_bom, "pending_checkout", return_value=("open", "https://checkout.stripe.test/one")), \
             patch.object(one_time_bom, "_service_client", return_value=client):
            with self.assertRaisesRegex(one_time_bom.OneTimeBOMError, "already open"):
                one_time_bom.begin_checkout(USER, "eng@example.com", "https://app/success", "https://app/cancel")
        client.table.assert_not_called()

    def test_stripe_checkout_is_one_payment_with_order_metadata(self):
        stripe = types.ModuleType("stripe")
        stripe.api_key = "sk_test_fake"
        create = Mock(return_value=types.SimpleNamespace(id="cs_test_one", url="https://checkout.stripe.test/one"))
        stripe.checkout = types.SimpleNamespace(Session=types.SimpleNamespace(create=create))
        helper_file = Path(__file__).resolve().parents[1] / "src/stripe_helper.py"
        spec = importlib.util.spec_from_file_location("cadivor_checkout_test", helper_file)
        helper = importlib.util.module_from_spec(spec)
        with patch.dict(sys.modules, {"stripe": stripe}):
            spec.loader.exec_module(helper)
            helper.create_one_time_bom_checkout(
                price_id=self.price_id, user_email="eng@example.com", user_id=USER,
                order_id=ORDER, success_url="https://app/success", cancel_url="https://app/cancel",
            )
        options = create.call_args.kwargs
        self.assertEqual(options["mode"], "payment")
        self.assertEqual(options["line_items"], [{"price": self.price_id, "quantity": 1}])
        self.assertEqual(options["metadata"]["cadivor_order_id"], ORDER)
        self.assertEqual(options["idempotency_key"], f"cadivor_single_bom_{ORDER}")

    def test_credit_is_reserved_and_consumed_only_for_saved_analysis(self):
        client = Mock()
        client.rpc.return_value.execute.return_value.data = ORDER
        with patch.object(one_time_bom, "_service_client", return_value=client), \
             patch.object(one_time_bom, "_orders_available", return_value=True):
            self.assertEqual(one_time_bom.reserve_credit(USER), ORDER)
            client.rpc.return_value.execute.return_value.data = True
            one_time_bom.attach_analysis(USER, ORDER, ANALYSIS)
            one_time_bom.consume_credit(USER, ORDER, ANALYSIS)
        self.assertEqual(client.rpc.call_args.args[0], "cadivor_consume_one_time_bom_order")
        self.assertEqual(client.rpc.call_args.args[1]["p_analysis_id"], ANALYSIS)

    def test_sales_kill_switch_does_not_revoke_a_paid_order(self):
        client = Mock()
        client.table.return_value.select.return_value.eq.return_value.eq.return_value.limit.return_value.execute.return_value.data = [{"id": ORDER}]
        with patch.object(one_time_bom, "_service_client", return_value=client), \
             patch.object(one_time_bom, "_orders_available", return_value=True), \
             patch.object(one_time_bom, "enabled", return_value=False):
            self.assertTrue(one_time_bom.available_credit(USER))

    def test_migration_keeps_entitlement_service_only_and_atomic(self):
        sql = (Path(__file__).resolve().parents[1] / "supabase/migrations/20261004_one_time_bom_reports.sql").read_text()
        self.assertIn("for update skip locked", sql)
        self.assertIn("v_order.stripe_session_id is distinct from p_session_id", sql)
        self.assertIn("v_order.price_id is distinct from p_price_id", sql)
        self.assertIn("and user_id = p_user_id", sql)
        self.assertIn("cadivor_attach_one_time_bom_analysis", sql)
        self.assertIn("and analysis_id = p_analysis_id", sql)
        self.assertIn("revoke all on public.cadivor_one_time_bom_orders from public, anon, authenticated", sql)
        self.assertIn("grant execute on function public.cadivor_fulfill_one_time_bom_order", sql)

    def test_full_bom_excel_is_generated_in_memory_per_request(self):
        report = BytesIO()
        save_results_to_excel([{
            "MPN": "MCP2551-I/SN", "Risk Score": 90, "Risk Level": "High",
            "Lifecycle Status": "Obsolete", "Stock Available": 0,
        }], report)
        workbook = load_workbook(BytesIO(report.getvalue()))
        self.assertEqual(workbook["Detailed Results"]["A2"].value, "MCP2551-I/SN")
        self.assertIn("High Risk Parts", workbook.sheetnames)


if __name__ == "__main__":
    unittest.main()
